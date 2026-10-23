"""
Excel 数据迁移与报表导出服务 (ExcelMigrationService)
支持：
1. 从内置 `default_seed.json` 或任意 `.xlsm/.xlsx` 原版表格一键导入关键字词库、盘口、分人统计客户、上庄名单与历史订单明细
2. 将当前盘口的汇总矩阵、1-49风险排行、分人算账对奖结果、订单明细与吃码上报导出为精美排版的 `.xlsx` 报表
"""

from __future__ import annotations
import json
import os
from pathlib import Path
import sys
from typing import Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from src.database.db_manager import DatabaseManager
from src.database.repositories import (
    ClientRepository,
    ConfigRepository,
    EatReportRepository,
    KeywordRepository,
    OrderRepository,
    RegionRepository,
)
from src.domain.attribute_engine import NumberAttributeEngine
from src.domain.models import (
    AppConfig,
    BookmakerRecord,
    ClientProfile,
    OrderRecord,
    RegionConfig,
)
from .risk_eat_service import RiskAndEatService
from .settlement_service import SettlementService


def get_resource_path(relative_path: str) -> Path:
    """获取打包进 PyInstaller exe 或开发源码环境下的资源文件路径（多重回退保护）"""
    rel = Path(relative_path)
    candidates: List[Path] = []
    if hasattr(sys, "_MEIPASS"):
        candidates.append(Path(getattr(sys, "_MEIPASS")) / rel)
    if getattr(sys, "frozen", False):
        candidates.append(Path(sys.executable).resolve().parent / rel)
    candidates.append(Path(__file__).resolve().parent.parent.parent / rel)
    candidates.append(Path.cwd() / rel)
    for c in candidates:
        if c.exists():
            return c
    return candidates[0]


class ExcelMigrationService:
    """面向对象 Excel 数据导入迁移与报表导出服务"""

    def __init__(
        self,
        db: DatabaseManager,
        config_repo: ConfigRepository,
        region_repo: RegionRepository,
        client_repo: ClientRepository,
        order_repo: OrderRepository,
        eat_repo: EatReportRepository,
        keyword_repo: KeywordRepository,
        attr_engine: NumberAttributeEngine,
        risk_service: RiskAndEatService,
        settlement_service: SettlementService,
    ) -> None:
        self.db = db
        self.config_repo = config_repo
        self.region_repo = region_repo
        self.client_repo = client_repo
        self.order_repo = order_repo
        self.eat_repo = eat_repo
        self.keyword_repo = keyword_repo
        self.attr_engine = attr_engine
        self.risk_service = risk_service
        self.settlement_service = settlement_service

    def initialize_if_needed(self, prefer_excel_path: Optional[str | Path] = None) -> bool:
        """
        如果本地 SQLite 数据库尚未初始化种子数据，则自动从 Excel 或内置 `default_seed.json` 迁移初始化
        """
        if self.db.is_initialized_with_seed_data():
            return False

        if prefer_excel_path and Path(prefer_excel_path).exists():
            try:
                self.import_from_excel_file(prefer_excel_path)
                return True
            except Exception:
                pass

        seed_path = get_resource_path("src/data/default_seed.json")
        if seed_path.exists():
            try:
                self.import_from_seed_json(seed_path)
                return True
            except Exception:
                pass

        # 终极兜底自愈：即使外部种子文件缺失，也自动写入核心关键字、盘口、默认客户与上庄，保证新用户 100% 正常使用
        self._initialize_minimal_fallback_seed()
        return True

    def _initialize_minimal_fallback_seed(self) -> None:
        """终极内置兜底初始化数据（零外部文件依赖）"""
        cfg = self.config_repo.load_app_config()
        self.config_repo.save_app_config(cfg)
        core_each_num = ["各号", "各码", "每个号", "每个码", "每号", "每码", "号各", "码各", "各买", "各下", "各"]
        core_each_zod = ["各肖", "每肖", "一肖", "每个肖", "肖各"]
        core_total = ["共", "一共", "总共", "合计", "平摊", "平均"]
        all_kws = list(dict.fromkeys(core_each_num + core_each_zod + core_total + ["米", "元", "块", "蚊"]))
        self.keyword_repo.bulk_insert_keywords(
            {
                "全部关键字": all_kws,
                "关键字-各肖": core_each_zod,
                "关键字-各号": core_each_num,
                "关键字-共": core_total,
            }
        )
        for idx, rname in enumerate(["新澳门", "香港", "旧澳门"], 1):
            self.region_repo.save_region(
                RegionConfig(name=rname, display_order=idx, global_odds=47.0, global_rebate=0.02, active_client_name="客户1")
            )
            for c_idx in range(1, 6):
                self.client_repo.save_client(
                    ClientProfile(id=None, region=rname, name=f"客户{c_idx}", odds=47.0, rebate=0.0, sort_order=c_idx)
                )
            for bm_name in ["上庄张", "上庄李", "上庄王", "上庄周", "上庄吴"]:
                self.eat_repo.save_bookmaker(BookmakerRecord(id=None, region=rname, name=bm_name))

    def import_from_seed_json(self, seed_json_path: str | Path) -> Dict[str, int]:
        """从内置 `default_seed.json` 导入全部关键字、盘口、客户、订单和上庄"""
        with open(seed_json_path, "r", encoding="utf-8") as f:
            seed = json.load(f)

        # 1. 保存默认全局配置
        cfg = self.config_repo.load_app_config()
        self.config_repo.save_app_config(cfg)

        # 2. 批量导入关键字
        kw_map: Dict[str, List[str]] = seed.get("keywords", {})
        self.keyword_repo.bulk_insert_keywords(kw_map)

        # 3. 导入盘口、客户、订单、上庄
        order_count = 0
        client_count = 0
        for reg_data in seed.get("regions", []):
            rname = reg_data["name"]
            reg_cfg = RegionConfig(
                name=rname,
                display_order=int(reg_data.get("display_order", 1)),
                global_odds=float(reg_data.get("global_odds", 47.0)),
                global_rebate=float(reg_data.get("global_rebate", 0.02)),
                active_client_name=reg_data.get("active_client_name", ""),
                last_winning_number=0,
            )
            self.region_repo.save_region(reg_cfg)

            for c_data in seed.get("clients", {}).get(rname, []):
                bets = {i: 0.0 for i in range(1, 50)}
                for k, v in c_data.get("bets", {}).items():
                    bets[int(k)] = float(v)
                cp = ClientProfile(
                    id=None,
                    region=rname,
                    name=c_data["name"],
                    odds=float(c_data.get("odds", 47.0)),
                    rebate=float(c_data.get("rebate", 0.0)),
                    sort_order=int(c_data.get("sort_order", 1)),
                    number_bets=bets,
                )
                self.client_repo.save_client(cp)
                client_count += 1

            self.order_repo.clear_region_orders(rname)
            for o_data in seed.get("orders", {}).get(rname, []):
                ord_rec = OrderRecord(
                    id=None,
                    seq_no=int(o_data["seq_no"]),
                    region=rname,
                    client_name=o_data["client_name"],
                    raw_text=o_data["raw_text"],
                    parsed_content=o_data["parsed_content"],
                    total_amount=float(o_data["total_amount"]),
                    created_at=o_data.get("created_at") or "",
                )
                self.order_repo.add_order(ord_rec)
                order_count += 1

            for bm_name in seed.get("bookmakers", {}).get(rname, ["上庄张", "上庄李", "上庄王", "上庄周", "上庄吴"]):
                self.eat_repo.save_bookmaker(BookmakerRecord(id=None, region=rname, name=bm_name))

        return {
            "keywords": sum(len(v) for v in kw_map.values()),
            "regions": len(seed.get("regions", [])),
            "clients": client_count,
            "orders": order_count,
        }

    def import_from_excel_file(self, excel_path: str | Path, clear_existing_orders: bool = True) -> Dict[str, int]:
        """直接读取原版 `.xlsm` / `.xlsx` 宏工作簿并迁移全部数据到本地 SQLite 数据库"""
        wb = openpyxl.load_workbook(str(excel_path), data_only=True)
        stats = {"keywords": 0, "regions": 0, "clients": 0, "orders": 0}

        # 1. 关键字列表
        if "关键字列表" in wb.sheetnames:
            ws_kw = wb["关键字列表"]
            cats = ["全部关键字", "关键字-各肖", "关键字-各号", "关键字-共"]
            kw_map: Dict[str, List[str]] = {c: [] for c in cats}
            for c_idx, cat in enumerate(cats, 1):
                for r in range(2, ws_kw.max_row + 1):
                    v = ws_kw.cell(r, c_idx).value
                    if v is not None and str(v).strip():
                        kw_map[cat].append(str(v).strip())
            self.keyword_repo.bulk_insert_keywords(kw_map)
            stats["keywords"] = sum(len(v) for v in kw_map.values())

        # 2. 盘口、分人统计、订单明细、数据明细
        candidate_regions = ["新澳门", "香港", "旧澳门"]
        for idx, rname in enumerate(candidate_regions, 1):
            if rname not in wb.sheetnames:
                continue
            ws_r = wb[rname]
            g_odds = float(ws_r["L14"].value or 47.0)
            g_rebate = float(ws_r["L15"].value or 0.02)
            active_c = str(ws_r["E1"].value or "").strip()

            reg_cfg = RegionConfig(
                name=rname,
                display_order=idx,
                global_odds=g_odds,
                global_rebate=g_rebate,
                active_client_name=active_c,
                last_winning_number=0,
            )
            self.region_repo.save_region(reg_cfg)
            stats["regions"] += 1

            # 客户分人统计
            stat_sheet_name = f"{rname}_分人统计"
            if stat_sheet_name in wb.sheetnames:
                ws_c = wb[stat_sheet_name]
                for r in range(1, ws_c.max_row + 1):
                    cname = ws_c.cell(r, 1).value
                    if not cname:
                        continue
                    cname = str(cname).strip()
                    bets = {i: 0.0 for i in range(1, 50)}
                    for num in range(1, 50):
                        val = ws_c.cell(r, num + 1).value
                        if val is not None and float(val) != 0:
                            bets[num] = round(float(val), 2)
                    c_odds = float(ws_c.cell(r, 51).value or g_odds)
                    c_rebate = float(ws_c.cell(r, 52).value or 0.0)
                    cp = ClientProfile(
                        id=None,
                        region=rname,
                        name=cname,
                        odds=c_odds,
                        rebate=c_rebate,
                        sort_order=r,
                        number_bets=bets,
                    )
                    self.client_repo.save_client(cp)
                    stats["clients"] += 1

            # 订单明细
            order_sheet_name = f"{rname}明细"
            if order_sheet_name in wb.sheetnames:
                ws_o = wb[order_sheet_name]
                if clear_existing_orders:
                    self.order_repo.clear_region_orders(rname)
                for r in range(3, ws_o.max_row + 1):
                    seq = ws_o.cell(r, 1).value
                    if seq is None:
                        continue
                    cname = str(ws_o.cell(r, 3).value or "").strip()
                    raw_t = str(ws_o.cell(r, 4).value or "").strip()
                    parsed_t = str(ws_o.cell(r, 5).value or "").strip()
                    amt = float(ws_o.cell(r, 6).value or 0.0)
                    dt_val = ws_o.cell(r, 11).value
                    dt_str = (
                        dt_val.strftime("%Y-%m-%d %H:%M:%S")
                        if hasattr(dt_val, "strftime")
                        else str(dt_val or "")
                    )
                    self.order_repo.add_order(
                        OrderRecord(
                            id=None,
                            seq_no=int(seq),
                            region=rname,
                            client_name=cname,
                            raw_text=raw_t,
                            parsed_content=parsed_t,
                            total_amount=amt,
                            created_at=dt_str,
                        )
                    )
                    stats["orders"] += 1

            # 上庄名单
            data_sheet_name = f"{rname}数据明细"
            if data_sheet_name in wb.sheetnames:
                ws_d = wb[data_sheet_name]
                for r in range(2, 10):
                    bname = ws_d.cell(r, 9).value
                    if bname and str(bname).strip():
                        self.eat_repo.save_bookmaker(
                            BookmakerRecord(id=None, region=rname, name=str(bname).strip())
                        )

        wb.close()
        return stats

    def export_region_report_to_excel(self, region: str, output_path: str | Path) -> Path:
        """
        将指定盘口的全盘统计、1-49风险榜、对奖算账结果、订单明细与吃码上报导出为精美排版的 `.xlsx` 文件
        """
        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        reg_cfg = self.region_repo.get_region(region) or RegionConfig(name=region)
        risk_items = self.risk_service.get_region_risk_items(region)
        sorted_risk = self.risk_service.get_sorted_risk_ranking(region)
        win_num = reg_cfg.last_winning_number if 1 <= reg_cfg.last_winning_number <= 49 else None
        client_settlements, grand_total = self.settlement_service.settle_region(region, win_num)
        orders = self.order_repo.list_orders(region)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="微软雅黑", size=10)
        red_font = Font(name="微软雅黑", size=10, bold=True, color="DC2626")
        green_font = Font(name="微软雅黑", size=10, bold=True, color="16A34A")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Sheet 1: 盘口汇总与风险排行
        ws1 = wb.active
        ws1.title = f"{region}_汇总与风险榜"
        ws1.append(
            [
                "排名", "号码", "生肖", "波色", "汇总金额",
                "自己吃码", "预备上报", "已经上报", "全盘风险值(盈亏)", "吃码风险值(盈亏)"
            ]
        )
        for col_idx in range(1, 11):
            cell = ws1.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        wave_cn = {"red": "红波", "blue": "蓝波", "green": "绿波"}
        for rank, it in enumerate(sorted_risk, start=1):
            ws1.append(
                [
                    rank,
                    it.number_str,
                    it.zodiac,
                    wave_cn.get(it.wave_color, ""),
                    it.total_bet,
                    it.eat_amount,
                    it.report_amount,
                    it.reported_amount,
                    it.full_risk_value,
                    it.eat_risk_value,
                ]
            )
            r_idx = rank + 1
            for c_idx in range(1, 11):
                c = ws1.cell(r_idx, c_idx)
                c.font = body_font
                c.border = thin_border
                c.alignment = center_align
            ws1.cell(r_idx, 9).font = red_font if it.full_risk_value < 0 else green_font
            ws1.cell(r_idx, 10).font = red_font if it.eat_risk_value < 0 else green_font

        # Sheet 2: 分人统计与对奖算账
        ws2 = wb.create_sheet(title=f"{region}_对奖算账")
        ws2.append(
            [
                "客户姓名", "订单笔数", "赔率", "回水比例",
                "下单总额", f"中码金额(特码:{win_num or '未开奖'})",
                "盈亏(不包回水)", "回水金额", "总盈亏(含回水)"
            ]
        )
        for col_idx in range(1, 10):
            cell = ws2.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        all_settle_rows = client_settlements + [grand_total]
        for idx, st in enumerate(all_settle_rows, start=2):
            ws2.append(
                [
                    st.client_name,
                    st.order_count,
                    st.odds if st.client_name != "合计：" else "",
                    st.rebate if st.client_name != "合计：" else "",
                    st.total_bet,
                    st.winning_bet,
                    st.pnl_no_rebate,
                    st.rebate_amount,
                    st.net_pnl,
                ]
            )
            for c_idx in range(1, 10):
                c = ws2.cell(idx, c_idx)
                c.font = body_font
                c.border = thin_border
                c.alignment = center_align
            ws2.cell(idx, 9).font = red_font if st.net_pnl < 0 else green_font

        # Sheet 3: 订单明细
        ws3 = wb.create_sheet(title=f"{region}_订单明细")
        ws3.append(
            [
                "序号", "地区", "下单人", "原始复制粘贴内容", "识别押注内容",
                "下单金额", "中码金额", "盈亏(不包回水)", "回水", "总盈亏(含回水)", "下单时间"
            ]
        )
        for col_idx in range(1, 12):
            cell = ws3.cell(1, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align

        for idx, o in enumerate(orders, start=2):
            ws3.append(
                [
                    o.seq_no,
                    o.region,
                    o.client_name,
                    o.raw_text,
                    o.parsed_content,
                    o.total_amount,
                    o.winning_amount,
                    o.pnl_no_rebate,
                    o.rebate_amount,
                    o.net_pnl,
                    o.created_at,
                ]
            )
            for c_idx in range(1, 12):
                c = ws3.cell(idx, c_idx)
                c.font = body_font
                c.border = thin_border
                c.alignment = center_align

        for ws in (ws1, ws2, ws3):
            for col in ws.columns:
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = 16

        wb.save(str(out))
        return out

    def load_from_default_seed_json(self) -> Dict[str, int]:
        """重新从内置 `default_seed.json` 加载默认词库与种子数据"""
        seed_path = get_resource_path("src/data/default_seed.json")
        if seed_path.exists():
            return self.import_from_seed_json(seed_path)
        return {}

    def import_from_excel_workbook(self, excel_path: str | Path) -> Dict[str, int]:
        """`import_from_excel_file` 的语义化别名"""
        return self.import_from_excel_file(excel_path)

    def export_to_excel(self, output_path: str | Path, region: Optional[str] = None) -> Path:
        """导出指定盘口（默认第一个或当前盘口）的完整 Excel 报表"""
        reg_name = region
        if not reg_name:
            regs = self.region_repo.list_regions()
            reg_name = regs[0].name if regs else "新澳门"
        return self.export_region_report_to_excel(reg_name, output_path)
