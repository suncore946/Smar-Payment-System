"""
按日期历史回看与期数归档服务 (HistoryReviewService)

核心能力：
1. 双数据源合并回看：同时检索「当前活动订单表 orders」与「已归档快照 daily_archives」，
   按订单自身的下单日期 (created_at) 过滤，并自动去重，保证「清空本盘」后历史仍可完整回看；
2. 区间重建：按日期区间重建 1-49 号押注汇总、风险值、分人算账、每日趋势与订单流水；
3. 模拟特码重算：支持指定任意 01~49 特码，对历史区间内全部订单重新对奖算账；
4. 归档 / 恢复 / 删除：清空本盘前一键归档当日账目快照，并可将任意归档恢复到活动盘口；
5. 导出：生成历史结账单文本与多 Sheet 精美 Excel 报表。
"""

from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.database.repositories import (
    ClientRepository,
    HistoryArchiveRepository,
    OrderRepository,
    RegionRepository,
)
from src.domain.attribute_engine import NumberAttributeEngine
from src.domain.models import (
    ClientProfile,
    ClientSettlementItem,
    DailyArchiveRecord,
    DailyTrendItem,
    DateRangeReviewSummary,
    OrderRecord,
    RiskItem,
)
from .risk_eat_service import RiskAndEatService
from .settlement_service import SettlementService
from .wechat_parser_service import WeChatParserService

ALL_REGIONS = "全部盘口"
ALL_CLIENTS = "全部客户"
DEFAULT_ODDS = 47.0
DEFAULT_REBATE = 0.02


@dataclass
class _ResolvedOrder:
    """内部使用：附带赔率 / 回水 / 来源信息的订单"""

    order: OrderRecord
    odds: float
    rebate: float
    source: str  # "live" | "archive"


class HistoryReviewService:
    """按日期历史回看、模拟重算、期数归档与恢复服务"""

    def __init__(
        self,
        order_repo: OrderRepository,
        archive_repo: HistoryArchiveRepository,
        client_repo: ClientRepository,
        region_repo: RegionRepository,
        settlement_service: SettlementService,
        risk_service: RiskAndEatService,
        attr_engine: NumberAttributeEngine,
    ) -> None:
        self.order_repo = order_repo
        self.archive_repo = archive_repo
        self.client_repo = client_repo
        self.region_repo = region_repo
        self.settlement_service = settlement_service
        self.risk_service = risk_service
        self.attr_engine = attr_engine

    # ------------------------------------------------------------------
    # 日期工具
    # ------------------------------------------------------------------
    @staticmethod
    def normalize_date(value: str | date | datetime | None) -> str:
        """统一转换为 YYYY-MM-DD 字符串；非法输入抛出 ValueError"""
        if value is None:
            return date.today().strftime("%Y-%m-%d")
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        text = str(value).strip().replace("/", "-").replace(".", "-")[:10]
        return datetime.strptime(text, "%Y-%m-%d").strftime("%Y-%m-%d")

    @staticmethod
    def quick_range(preset: str, today: Optional[date] = None) -> Tuple[str, str]:
        """快捷日期区间：今天 / 昨天 / 近7天 / 近30天 / 本月 / 上月"""
        t = today or date.today()
        if preset == "今天":
            s = e = t
        elif preset == "昨天":
            s = e = t - timedelta(days=1)
        elif preset == "近7天":
            s, e = t - timedelta(days=6), t
        elif preset == "近30天":
            s, e = t - timedelta(days=29), t
        elif preset == "本月":
            s, e = t.replace(day=1), t
        elif preset == "上月":
            e = t.replace(day=1) - timedelta(days=1)
            s = e.replace(day=1)
        else:
            raise ValueError(f"未知的快捷区间：{preset}")
        return s.strftime("%Y-%m-%d"), e.strftime("%Y-%m-%d")

    # ------------------------------------------------------------------
    # 可用日期
    # ------------------------------------------------------------------
    def get_all_available_dates(self, region: Optional[str] = None) -> List[str]:
        """返回所有存在订单（活动表或归档内订单）的日期，降序排列"""
        dates = set(self.order_repo.list_distinct_order_dates(region))
        for arc in self.archive_repo.list_archives(region=region):
            dates.add(arc.archive_date)
            for od in arc.orders_data:
                created = str(od.get("created_at", ""))[:10]
                if created:
                    dates.add(created)
        return sorted((d for d in dates if d), reverse=True)

    # ------------------------------------------------------------------
    # 赔率 / 回水解析
    # ------------------------------------------------------------------
    def _live_odds_rebate(self, region: str, client_name: str) -> Tuple[float, float]:
        return self.settlement_service._get_client_odds_and_rebate(region, client_name)

    @staticmethod
    def _dedup_key(o: OrderRecord) -> Tuple:
        return (
            o.created_at,
            o.region,
            o.client_name,
            (o.parsed_content or "").strip(),
            round(float(o.total_amount), 2),
        )

    @staticmethod
    def _order_from_dict(d: Dict) -> OrderRecord:
        return OrderRecord(
            id=None,
            seq_no=int(d.get("seq_no", 0) or 0),
            region=str(d.get("region", "")),
            client_name=str(d.get("client_name", "")),
            raw_text=str(d.get("raw_text", "")),
            parsed_content=str(d.get("parsed_content", "")),
            total_amount=float(d.get("total_amount", 0.0) or 0.0),
            winning_amount=float(d.get("winning_amount", 0.0) or 0.0),
            pnl_no_rebate=float(d.get("pnl_no_rebate", 0.0) or 0.0),
            rebate_amount=float(d.get("rebate_amount", 0.0) or 0.0),
            net_pnl=float(d.get("net_pnl", 0.0) or 0.0),
            created_at=str(d.get("created_at", "")),
        )

    # ------------------------------------------------------------------
    # 订单收集（活动表 + 归档合并去重）
    # ------------------------------------------------------------------
    def collect_orders(
        self,
        start_date: str,
        end_date: str,
        region: Optional[str] = None,
        client_filter: Optional[str] = None,
        keyword_filter: Optional[str] = None,
        include_archives: bool = True,
    ) -> List[_ResolvedOrder]:
        s = self.normalize_date(start_date)
        e = self.normalize_date(end_date)
        if s > e:
            s, e = e, s

        resolved: List[_ResolvedOrder] = []
        seen: set = set()

        for o in self.order_repo.list_orders_by_date_range(s, e, region, client_filter, keyword_filter):
            key = self._dedup_key(o)
            if key in seen:
                continue
            seen.add(key)
            odds, rebate = self._live_odds_rebate(o.region, o.client_name)
            resolved.append(_ResolvedOrder(o, odds, rebate, "live"))

        if include_archives:
            kw = (keyword_filter or "").strip()
            use_client = client_filter if client_filter and client_filter not in ("全部", ALL_CLIENTS) else None
            # 归档日期可能被用户自定义，因此扫描该盘口全部归档，再按订单自身下单日期精确过滤
            for arc in self.archive_repo.list_archives(region=region):
                odds_map = {
                    str(st.get("client_name")): (
                        float(st.get("odds", DEFAULT_ODDS) or DEFAULT_ODDS),
                        float(st.get("rebate", 0.0) or 0.0),
                    )
                    for st in arc.settlements_data
                }
                for od in arc.orders_data:
                    o = self._order_from_dict(od)
                    if not o.region:
                        o.region = arc.region
                    d = o.created_at[:10]
                    if not d or d < s or d > e:
                        continue
                    if use_client and o.client_name != use_client:
                        continue
                    if kw and kw not in o.raw_text and kw not in o.parsed_content and kw not in o.client_name:
                        continue
                    key = self._dedup_key(o)
                    if key in seen:
                        continue
                    seen.add(key)
                    odds, rebate = odds_map.get(o.client_name) or self._live_odds_rebate(o.region, o.client_name)
                    resolved.append(_ResolvedOrder(o, odds, rebate, "archive"))

        resolved.sort(key=lambda r: (r.order.created_at, r.order.region, r.order.seq_no))
        return resolved

    # ------------------------------------------------------------------
    # 区间回看主入口
    # ------------------------------------------------------------------
    @staticmethod
    def _recalc(order: OrderRecord, win_num: int, odds: float, rebate: float) -> OrderRecord:
        """按模拟特码对单笔订单重新对奖（返回新对象，不修改数据库）"""
        num_map = WeChatParserService.parse_canonical_order_content(order.parsed_content)
        win_amt = round(num_map.get(win_num, 0.0), 2)
        pnl_no_reb = round(order.total_amount - odds * win_amt, 2)
        reb_amt = round(-order.total_amount * rebate, 2)
        new = OrderRecord(**asdict(order))
        new.winning_amount = win_amt
        new.pnl_no_rebate = pnl_no_reb
        new.rebate_amount = reb_amt
        new.net_pnl = round(pnl_no_reb + reb_amt, 2)
        return new

    def query_date_range_review(
        self,
        start_date: str,
        end_date: str,
        region: Optional[str] = None,
        client_filter: Optional[str] = None,
        keyword_filter: Optional[str] = None,
        simulated_winning_number: Optional[int] = None,
        include_archives: bool = True,
    ) -> DateRangeReviewSummary:
        s = self.normalize_date(start_date)
        e = self.normalize_date(end_date)
        if s > e:
            s, e = e, s
        sim = int(simulated_winning_number or 0)
        use_sim = 1 <= sim <= 49
        region_label = region if region else ALL_REGIONS

        resolved = self.collect_orders(s, e, region, client_filter, keyword_filter, include_archives)
        multi_region = len({r.order.region for r in resolved}) > 1

        number_bets: Dict[int, float] = {i: 0.0 for i in range(1, 50)}
        client_stats: Dict[str, ClientSettlementItem] = {}
        day_stats: Dict[str, Dict] = {}
        final_orders: List[OrderRecord] = []
        tot_bet = tot_payout = tot_reb = tot_net = 0.0

        for r in resolved:
            o = self._recalc(r.order, sim, r.odds, r.rebate) if use_sim else r.order
            final_orders.append(o)
            payout = round(r.odds * o.winning_amount, 2)

            num_map = WeChatParserService.parse_canonical_order_content(o.parsed_content)
            for num, amt in num_map.items():
                if 1 <= num <= 49:
                    number_bets[num] = round(number_bets[num] + amt, 2)

            c_key = f"{o.region}·{o.client_name}" if multi_region else o.client_name
            st = client_stats.get(c_key)
            if st is None:
                st = ClientSettlementItem(
                    client_name=c_key, odds=r.odds, rebate=r.rebate, total_bet=0.0,
                    winning_bet=0.0, pnl_no_rebate=0.0, rebate_amount=0.0, net_pnl=0.0, order_count=0,
                )
                client_stats[c_key] = st
            st.total_bet = round(st.total_bet + o.total_amount, 2)
            st.winning_bet = round(st.winning_bet + o.winning_amount, 2)
            st.pnl_no_rebate = round(st.pnl_no_rebate + o.pnl_no_rebate, 2)
            st.rebate_amount = round(st.rebate_amount + o.rebate_amount, 2)
            st.net_pnl = round(st.net_pnl + o.net_pnl, 2)
            st.order_count += 1

            d = o.created_at[:10]
            ds = day_stats.setdefault(
                d, {"orders": 0, "clients": set(), "bet": 0.0, "payout": 0.0, "reb": 0.0, "net": 0.0,
                    "nums": {}},
            )
            ds["orders"] += 1
            ds["clients"].add((o.region, o.client_name))
            ds["bet"] = round(ds["bet"] + o.total_amount, 2)
            ds["payout"] = round(ds["payout"] + payout, 2)
            ds["reb"] = round(ds["reb"] + o.rebate_amount, 2)
            ds["net"] = round(ds["net"] + o.net_pnl, 2)
            for num, amt in num_map.items():
                if 1 <= num <= 49:
                    ds["nums"][num] = round(ds["nums"].get(num, 0.0) + amt, 2)

            tot_bet = round(tot_bet + o.total_amount, 2)
            tot_payout = round(tot_payout + payout, 2)
            tot_reb = round(tot_reb + o.rebate_amount, 2)
            tot_net = round(tot_net + o.net_pnl, 2)

        trends: List[DailyTrendItem] = []
        for d, ds in day_stats.items():
            top_n, top_v = self._top_of(ds["nums"])
            trends.append(
                DailyTrendItem(
                    date_str=d, order_count=ds["orders"], client_count=len(ds["clients"]),
                    total_bet=ds["bet"], winning_payout=ds["payout"], rebate_amount=ds["reb"],
                    net_pnl=ds["net"], top_number=top_n, top_number_bet=top_v,
                )
            )
        trends.sort(key=lambda t: t.date_str, reverse=True)

        top_n, top_v = self._top_of(number_bets)
        risk_items = self._build_risk_items(number_bets, region)
        all_clients = {(o.region, o.client_name) for o in final_orders}

        return DateRangeReviewSummary(
            start_date=s,
            end_date=e,
            region=region_label,
            simulated_winning_number=sim if use_sim else 0,
            total_orders=len(final_orders),
            active_days_count=len(day_stats),
            active_clients_count=len(all_clients),
            total_bet=tot_bet,
            total_winning_payout=tot_payout,
            total_rebate=tot_reb,
            total_net_pnl=tot_net,
            top_number=top_n,
            top_number_bet=top_v,
            number_bets=number_bets,
            risk_items=risk_items,
            client_settlements=sorted(client_stats.values(), key=lambda c: -c.total_bet),
            orders=final_orders,
            daily_trends=trends,
        )

    @staticmethod
    def _top_of(nums: Dict[int, float]) -> Tuple[int, float]:
        best_n, best_v = 0, 0.0
        for n in sorted(nums):
            v = nums[n]
            if v > best_v:
                best_n, best_v = n, v
        return best_n, round(best_v, 2)

    def _build_risk_items(self, number_bets: Dict[int, float], region: Optional[str]) -> List[RiskItem]:
        reg_cfg = self.region_repo.get_region(region) if region and region != ALL_REGIONS else None
        odds = reg_cfg.global_odds if reg_cfg else DEFAULT_ODDS
        rebate = reg_cfg.global_rebate if reg_cfg else DEFAULT_REBATE
        pool = round(sum(number_bets.values()), 2)
        items: List[RiskItem] = []
        for num in range(1, 50):
            tot = round(number_bets.get(num, 0.0), 2)
            items.append(
                RiskItem(
                    number=num,
                    zodiac=self.attr_engine.get_zodiac_of_number(num),
                    wave_color=self.attr_engine.get_wave_color(num),
                    total_bet=tot,
                    full_risk_value=RiskAndEatService.calculate_risk_value(pool, tot, odds, rebate),
                )
            )
        return items

    # ------------------------------------------------------------------
    # 归档 / 恢复 / 删除
    # ------------------------------------------------------------------
    def archive_current_region_session(
        self,
        region: str,
        archive_date: Optional[str] = None,
        custom_title: str = "",
    ) -> Optional[DailyArchiveRecord]:
        """将指定盘口当前全部订单、分人算账与风险榜打包为一条历史归档快照；无订单时返回 None"""
        orders = self.order_repo.list_orders(region)
        if not orders:
            return None

        a_date = self.normalize_date(archive_date) if archive_date else max(o.created_at[:10] for o in orders)
        reg_cfg = self.region_repo.get_region(region)
        win_num = reg_cfg.last_winning_number if reg_cfg and 1 <= reg_cfg.last_winning_number <= 49 else 0

        settlements: Dict[str, Dict] = {}
        tot_bet = tot_win_payout = tot_reb = tot_net = 0.0
        for o in orders:
            odds, rebate = self._live_odds_rebate(region, o.client_name)
            st = settlements.setdefault(
                o.client_name,
                {"client_name": o.client_name, "odds": odds, "rebate": rebate, "total_bet": 0.0,
                 "winning_bet": 0.0, "pnl_no_rebate": 0.0, "rebate_amount": 0.0, "net_pnl": 0.0,
                 "order_count": 0},
            )
            st["total_bet"] = round(st["total_bet"] + o.total_amount, 2)
            st["winning_bet"] = round(st["winning_bet"] + o.winning_amount, 2)
            st["pnl_no_rebate"] = round(st["pnl_no_rebate"] + o.pnl_no_rebate, 2)
            st["rebate_amount"] = round(st["rebate_amount"] + o.rebate_amount, 2)
            st["net_pnl"] = round(st["net_pnl"] + o.net_pnl, 2)
            st["order_count"] += 1
            tot_bet = round(tot_bet + o.total_amount, 2)
            tot_win_payout = round(tot_win_payout + odds * o.winning_amount, 2)
            tot_reb = round(tot_reb + o.rebate_amount, 2)
            tot_net = round(tot_net + o.net_pnl, 2)

        risk_data = [asdict(it) for it in self.risk_service.get_region_risk_items(region) if it.total_bet > 0]
        existing = self.archive_repo.list_archives(start_date=a_date, end_date=a_date, region=region)
        title = custom_title.strip() if custom_title and custom_title.strip() else f"{a_date} {region} 第{len(existing) + 1}期归档"

        record = DailyArchiveRecord(
            id=None,
            archive_date=a_date,
            region=region,
            title=title,
            winning_number=win_num,
            base_zodiac=self.attr_engine.base_zodiac,
            total_bet=tot_bet,
            total_winning=tot_win_payout,
            total_rebate=tot_reb,
            net_pnl=tot_net,
            order_count=len(orders),
            client_count=len(settlements),
            orders_data=[asdict(o) for o in orders],
            settlements_data=list(settlements.values()),
            risk_data=risk_data,
        )
        self.archive_repo.save_archive(record)
        return record

    def list_archives(
        self, start_date: Optional[str] = None, end_date: Optional[str] = None, region: Optional[str] = None
    ) -> List[DailyArchiveRecord]:
        return self.archive_repo.list_archives(start_date, end_date, region)

    def delete_archive(self, archive_id: int) -> bool:
        if not self.archive_repo.get_archive_by_id(archive_id):
            return False
        self.archive_repo.delete_archive(archive_id)
        return True

    def restore_archive_to_region(self, archive_id: int, target_region: Optional[str] = None) -> int:
        """
        将归档中的订单恢复到活动盘口（保留原下单时间），已存在的相同订单自动跳过；
        恢复后重建客户 1-49 号押注汇总。返回实际恢复的订单数。
        """
        arc = self.archive_repo.get_archive_by_id(archive_id)
        if not arc:
            raise ValueError(f"归档记录不存在：#{archive_id}")
        region = target_region or arc.region
        if not self.region_repo.get_region(region):
            from src.domain.models import RegionConfig

            self.region_repo.save_region(RegionConfig(name=region, display_order=99))

        existing_keys = {self._dedup_key(o) for o in self.order_repo.list_orders(region)}
        settle_map = {str(s.get("client_name")): s for s in arc.settlements_data}
        restored = 0
        for od in arc.orders_data:
            o = self._order_from_dict(od)
            o.region = region
            o.seq_no = 0
            if self._dedup_key(o) in existing_keys:
                continue
            if not self.client_repo.get_client(region, o.client_name):
                st = settle_map.get(o.client_name, {})
                sort_order = len(self.client_repo.list_clients(region)) + 1
                self.client_repo.save_client(
                    ClientProfile(
                        id=None, region=region, name=o.client_name,
                        odds=float(st.get("odds", DEFAULT_ODDS) or DEFAULT_ODDS),
                        rebate=float(st.get("rebate", 0.0) or 0.0),
                        sort_order=sort_order,
                    )
                )
            self.order_repo.add_order(o)
            existing_keys.add(self._dedup_key(o))
            restored += 1

        if restored:
            self.settlement_service.rebuild_region_client_bets_from_orders(region)
        return restored

    # ------------------------------------------------------------------
    # 文本结账单 / Excel 导出
    # ------------------------------------------------------------------
    @staticmethod
    def _date_label(summary: DateRangeReviewSummary) -> str:
        if summary.start_date == summary.end_date:
            return summary.start_date
        return f"{summary.start_date} ~ {summary.end_date}"

    def format_history_settlement_bill(self, summary: DateRangeReviewSummary) -> str:
        lines = [
            f"【历史结账单】{self._date_label(summary)}  盘口：{summary.region}",
        ]
        if summary.simulated_winning_number:
            lines.append(f"模拟特码：{summary.simulated_winning_number:02d}（按模拟号码重算）")
        else:
            lines.append("对奖口径：按原单记录")
        lines.append(
            f"订单 {summary.total_orders} 笔 / 活跃 {summary.active_days_count} 天 / 客户 {summary.active_clients_count} 人"
        )
        lines.append("-" * 28)
        for st in summary.client_settlements:
            lines.append(
                f"{st.client_name}：下单 {st.total_bet:g}，中码 {st.winning_bet:g}，"
                f"回水 {st.rebate_amount:g}，盈亏 {st.net_pnl:g}"
            )
        lines.append("-" * 28)
        lines.append(
            f"合计：下单 {summary.total_bet:g}，赔付 {summary.total_winning_payout:g}，"
            f"回水 {summary.total_rebate:g}，庄家净盈亏 {summary.total_net_pnl:g}"
        )
        if summary.top_number:
            lines.append(f"区间头号热门码：{summary.top_number:02d}（{summary.top_number_bet:g}）")
        return "\n".join(lines)

    def export_history_review_to_excel(self, summary: DateRangeReviewSummary, file_path: str | Path) -> Path:
        out = Path(file_path)
        out.parent.mkdir(parents=True, exist_ok=True)

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="微软雅黑", size=10)
        red_font = Font(name="微软雅黑", size=10, bold=True, color="DC2626")
        green_font = Font(name="微软雅黑", size=10, bold=True, color="16A34A")
        border = Border(*(Side(style="thin", color="CBD5E1") for _ in range(4)))
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def write_sheet(ws, headers: List[str], rows: Iterable[List], pnl_cols: Iterable[int] = ()) -> None:
            ws.append(headers)
            for c in range(1, len(headers) + 1):
                cell = ws.cell(1, c)
                cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
                ws.column_dimensions[cell.column_letter].width = 16
            pnl_set = set(pnl_cols)
            for r_idx, row in enumerate(rows, start=2):
                ws.append(row)
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(r_idx, c)
                    cell.font, cell.alignment, cell.border = body_font, center, border
                    if c in pnl_set and isinstance(cell.value, (int, float)):
                        cell.font = green_font if cell.value >= 0 else red_font
            ws.freeze_panes = "A2"

        wb = openpyxl.Workbook()
        ws0 = wb.active
        ws0.title = "区间汇总"
        write_sheet(
            ws0,
            ["项目", "数值"],
            [
                ["日期区间", self._date_label(summary)],
                ["盘口", summary.region],
                ["对奖口径", f"模拟特码 {summary.simulated_winning_number:02d}" if summary.simulated_winning_number else "按原单记录"],
                ["订单笔数", summary.total_orders],
                ["活跃天数", summary.active_days_count],
                ["活跃客户", summary.active_clients_count],
                ["区间总收单额", summary.total_bet],
                ["区间中码赔付", summary.total_winning_payout],
                ["区间回水", summary.total_rebate],
                ["庄家净盈亏", summary.total_net_pnl],
                ["头号热门码", f"{summary.top_number:02d}" if summary.top_number else "-"],
                ["热门码押注额", summary.top_number_bet],
            ],
        )

        write_sheet(
            wb.create_sheet("按日趋势"),
            ["日期", "订单数", "客户数", "收单额", "中码赔付", "回水", "庄家净盈亏", "热门码", "热门码押注"],
            [
                [t.date_str, t.order_count, t.client_count, t.total_bet, t.winning_payout, t.rebate_amount,
                 t.net_pnl, f"{t.top_number:02d}" if t.top_number else "-", t.top_number_bet]
                for t in summary.daily_trends
            ],
            pnl_cols=[7],
        )
        write_sheet(
            wb.create_sheet("分人算账"),
            ["客户", "赔率", "回水率", "订单数", "下单总额", "中码金额", "盈亏(不含回水)", "回水金额", "总盈亏"],
            [
                [c.client_name, c.odds, c.rebate, c.order_count, c.total_bet, c.winning_bet, c.pnl_no_rebate,
                 c.rebate_amount, c.net_pnl]
                for c in summary.client_settlements
            ],
            pnl_cols=[7, 9],
        )
        wave_cn = {"red": "红波", "blue": "蓝波", "green": "绿波"}
        write_sheet(
            wb.create_sheet("1-49号码复盘"),
            ["号码", "生肖", "波色", "区间押注额", "全盘风险值(盈亏)"],
            [[it.number_str, it.zodiac, wave_cn.get(it.wave_color, ""), it.total_bet, it.full_risk_value]
             for it in summary.risk_items],
            pnl_cols=[5],
        )
        write_sheet(
            wb.create_sheet("订单流水"),
            ["下单时间", "盘口", "客户", "原始内容", "识别内容", "下单金额", "中码金额", "回水", "总盈亏"],
            [
                [o.created_at, o.region, o.client_name, o.raw_text, o.parsed_content, o.total_amount,
                 o.winning_amount, o.rebate_amount, o.net_pnl]
                for o in summary.orders
            ],
            pnl_cols=[9],
        )
        wb.save(str(out))
        return out
