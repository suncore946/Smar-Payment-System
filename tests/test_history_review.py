"""
按日期历史回看与期数归档服务测试 (test_history_review.py)

测试场景涵盖：
1. 多日期模拟订单写入与按单日/跨日区间精准过滤查询；
2. 历史区间 1-49 号码汇总、热门号码统计与每日趋势指标计算；
3. 模拟特码动态重新对奖算账（只读推演，不篡改原单据）；
4. 当期盘口账目归档快照保存、清空盘口后历史仍完整可查、以及快照一键恢复至活动盘口；
5. 历史区间分人结账单文本生成与多 Sheet Excel 文件导出校验；
6. HistoryReviewTab UI 控件交互、快捷区间切换与归档恢复端到端测试。
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
import openpyxl
import pytest

from src.domain.models import OrderRecord, ParsedBetItem
from src.services.app_controller import ApplicationController
from src.ui.history_review_tab import HistoryReviewTab
from src.ui.main_window import MainWindow


class TestHistoryReviewService:
    """历史回看与归档服务核心逻辑测试"""

    def test_multi_date_orders_and_range_query(self, controller: ApplicationController) -> None:
        """测试不同日期订单的录入与单日/跨日区间查询"""
        region = controller.current_region
        c1 = controller.get_active_client(region).name

        # 模拟 3 天不同的订单数据
        # Day 1: 2026-10-20
        o1 = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="01,02各100",
            parsed_content="01,02各100",
            total_amount=200.0,
            created_at="2026-10-20 10:00:00",
        )
        # Day 2: 2026-10-21
        o2 = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="02,03各200",
            parsed_content="02,03各200",
            total_amount=400.0,
            created_at="2026-10-21 15:30:00",
        )
        # Day 3: 2026-10-22
        o3 = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="01各300",
            parsed_content="01各300",
            total_amount=300.0,
            created_at="2026-10-22 20:00:00",
        )

        controller.order_repo.add_order(o1)
        controller.order_repo.add_order(o2)
        controller.order_repo.add_order(o3)

        # 1. 验证可用日期列表
        available_dates = controller.history_service.get_all_available_dates(region)
        assert "2026-10-20" in available_dates
        assert "2026-10-21" in available_dates
        assert "2026-10-22" in available_dates

        # 2. 查询单日 2026-10-20
        sum_day1 = controller.history_service.query_date_range_review(
            start_date="2026-10-20",
            end_date="2026-10-20",
            region=region,
        )
        assert sum_day1.total_orders == 1
        assert sum_day1.total_bet == 200.0
        assert sum_day1.active_days_count == 1
        assert sum_day1.number_bets[1] == 100.0
        assert sum_day1.number_bets[2] == 100.0
        assert sum_day1.number_bets[3] == 0.0

        # 3. 跨日区间查询 2026-10-20 ~ 2026-10-22
        sum_all = controller.history_service.query_date_range_review(
            start_date="2026-10-20",
            end_date="2026-10-22",
            region=region,
        )
        assert sum_all.total_orders == 3
        assert sum_all.total_bet == 900.0
        assert sum_all.active_days_count == 3
        # 号码累计：01 号 = 100 + 300 = 400; 02 号 = 100 + 200 = 300; 03 号 = 200
        assert sum_all.number_bets[1] == 400.0
        assert sum_all.number_bets[2] == 300.0
        assert sum_all.number_bets[3] == 200.0
        # 头号热门码应当是 01 号 (400)
        assert sum_all.top_number == 1
        assert sum_all.top_number_bet == 400.0

        # 4. 验证每日趋势
        assert len(sum_all.daily_trends) == 3
        trend_dates = [t.date_str for t in sum_all.daily_trends]
        assert trend_dates == ["2026-10-22", "2026-10-21", "2026-10-20"]
        t21 = next(t for t in sum_all.daily_trends if t.date_str == "2026-10-21")
        assert t21.order_count == 1
        assert t21.total_bet == 400.0
        assert t21.top_number == 2
        assert t21.top_number_bet == 200.0

    def test_simulated_winning_number_recalculation(self, clean_controller: ApplicationController) -> None:
        """测试指定模拟特码对历史单据重新动态对奖"""
        controller = clean_controller
        region = controller.current_region
        c1 = controller.get_active_client(region).name
        # 预设赔率 47.0，回水 0.02
        client_obj = controller.client_repo.get_client(region, c1)
        client_obj.rebate = 0.02
        controller.client_repo.save_client(client_obj)

        order = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="08各100",
            parsed_content="08各100",
            total_amount=100.0,
            created_at="2026-10-21 12:00:00",
        )
        controller.order_repo.add_order(order)

        # 1. 模拟特码为 08（中码）：中码额=100，赔付=4700，回水=-2，净盈亏=100 - 4700 - 2 = -4602
        sum_hit = controller.history_service.query_date_range_review(
            start_date="2026-10-21",
            end_date="2026-10-21",
            region=region,
            simulated_winning_number=8,
        )
        assert sum_hit.simulated_winning_number == 8
        assert sum_hit.total_winning_payout == 4700.0
        assert sum_hit.total_net_pnl == -4602.0
        assert sum_hit.orders[0].winning_amount == 100.0
        assert sum_hit.orders[0].net_pnl == -4602.0

        # 2. 模拟特码为 09（未中码）：中码额=0，赔付=0，回水=-2，净盈亏=100 - 2 = +98
        sum_miss = controller.history_service.query_date_range_review(
            start_date="2026-10-21",
            end_date="2026-10-21",
            region=region,
            simulated_winning_number=9,
        )
        assert sum_miss.total_winning_payout == 0.0
        assert sum_miss.total_net_pnl == 98.0
        assert sum_miss.orders[0].winning_amount == 0.0
        assert sum_miss.orders[0].net_pnl == 98.0

        # 3. 验证原数据库中的单据并没有被破坏或篡改
        db_order = controller.order_repo.get_order_by_id(order.id)
        assert db_order.winning_amount == 0.0

    def test_archive_clear_and_restore_workflow(self, clean_controller: ApplicationController) -> None:
        """测试归档、清空盘口、历史回看依然保留、以及归档一键恢复"""
        controller = clean_controller
        region = controller.current_region
        c1 = controller.get_active_client(region).name

        order = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="05,06各50",
            parsed_content="05,06各50",
            total_amount=100.0,
            created_at="2026-10-21 16:00:00",
        )
        controller.order_repo.add_order(order)
        controller.client_repo.add_client_bets(region, c1, {5: 50.0, 6: 50.0})

        # 1. 归档当前会话
        archive = controller.history_service.archive_current_region_session(
            region=region,
            archive_date="2026-10-21",
            custom_title="2026-10-21 测试归档",
        )
        assert archive is not None
        assert archive.id is not None
        assert archive.order_count == 1
        assert archive.total_bet == 100.0

        # 2. 清空本盘
        controller.settlement_service.clear_region_all_data(region)
        assert len(controller.order_repo.list_orders(region)) == 0
        assert controller.client_repo.get_region_total_bets_by_number(region)[5] == 0.0

        # 3. 验证历史回看服务依然能从 daily_archives 中合并回看该日期的订单！
        review = controller.history_service.query_date_range_review(
            start_date="2026-10-21",
            end_date="2026-10-21",
            region=region,
        )
        assert review.total_orders == 1
        assert review.total_bet == 100.0
        assert review.number_bets[5] == 50.0
        assert review.number_bets[6] == 50.0

        # 4. 将归档恢复回活动盘口
        restored_cnt = controller.history_service.restore_archive_to_region(archive.id, region)
        assert restored_cnt == 1

        active_orders = controller.order_repo.list_orders(region)
        assert len(active_orders) == 1
        assert active_orders[0].parsed_content == "05,06各50"

        # 校验客户 1-49 押注已被自动同步重建
        total_bets = controller.client_repo.get_region_total_bets_by_number(region)
        assert total_bets[5] == 50.0
        assert total_bets[6] == 50.0

    def test_bill_generation_and_excel_export(self, tmp_path: Path, controller: ApplicationController) -> None:
        """测试结账单文本格式化与多 Sheet Excel 报表导出"""
        region = controller.current_region
        c1 = controller.get_active_client(region).name

        order = OrderRecord(
            id=None,
            seq_no=0,
            region=region,
            client_name=c1,
            raw_text="09各100",
            parsed_content="09各100",
            total_amount=100.0,
            created_at="2026-10-22 18:00:00",
        )
        controller.order_repo.add_order(order)

        summary = controller.history_service.query_date_range_review(
            start_date="2026-10-22",
            end_date="2026-10-22",
            region=region,
            simulated_winning_number=9,
        )

        # 1. 验证结账单文本
        bill_txt = controller.history_service.format_history_settlement_bill(summary)
        assert "【历史结账单】" in bill_txt
        assert "2026-10-22" in bill_txt
        assert "模拟特码：09" in bill_txt
        assert c1 in bill_txt

        # 2. 导出 Excel
        export_file = tmp_path / "test_history_report.xlsx"
        res_path = controller.history_service.export_history_review_to_excel(summary, export_file)
        assert res_path.exists()

        wb = openpyxl.load_workbook(str(res_path))
        sheet_names = wb.sheetnames
        assert "区间汇总" in sheet_names
        assert "按日趋势" in sheet_names
        assert "分人算账" in sheet_names
        assert "1-49号码复盘" in sheet_names
        assert "订单流水" in sheet_names

        ws_sum = wb["区间汇总"]
        assert ws_sum.cell(2, 1).value == "日期区间"
        assert ws_sum.cell(2, 2).value == "2026-10-22"


class TestHistoryReviewUI:
    """历史回看 UI 标签页与主界面集成测试"""

    def test_history_tab_ui_interaction(self, qapp, controller: ApplicationController) -> None:
        """测试 HistoryReviewTab 的快捷按钮、查询、归档与快照表格联动"""
        window = MainWindow(controller)
        window.show()
        qapp.processEvents()

        hist_tab = window.history_review_tab

        # 切换到历史回看页 (Alt+6 或直接切换控件)
        assert window.trigger_shortcut_action("switch_tab_history") is True
        assert window.tabs.currentWidget() == hist_tab

        # 快捷日期区间测试
        hist_tab.set_quick_range("近7天")
        qapp.processEvents()
        today = date.today()
        seven_days_ago = today - timedelta(days=6)
        assert hist_tab.date_start.date().toString("yyyy-MM-dd") == seven_days_ago.strftime("%Y-%m-%d")
        assert hist_tab.date_end.date().toString("yyyy-MM-dd") == today.strftime("%Y-%m-%d")

        # 录入一笔测试订单后归档
        region = controller.current_region
        c1 = controller.get_active_client(region).name
        controller.order_repo.add_order(
            OrderRecord(
                id=None,
                seq_no=0,
                region=region,
                client_name=c1,
                raw_text="12各50",
                parsed_content="12各50",
                total_amount=50.0,
                created_at=today.strftime("%Y-%m-%d 10:00:00"),
            )
        )

        hist_tab.set_quick_range("今天")
        summary = hist_tab.run_query()
        qapp.processEvents()
        assert summary.total_orders >= 1
        assert hist_tab.order_table.rowCount() >= 1

        # 测试在 UI 上执行归档
        rec = hist_tab.archive_current("UI测试自动化归档")
        qapp.processEvents()
        assert rec is not None
        assert hist_tab.archive_table.rowCount() >= 1

        # 测试剪贴板复制结账单文本
        bill = hist_tab.build_bill_text()
        assert "【历史结账单】" in bill

        window.close()
